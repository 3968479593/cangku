"""选定商品批量 AI 分析 + 导出 Word 报告。"""
import io
import json
import logging
from datetime import datetime

from fastapi import APIRouter
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from app.config import settings
from app.core.llm import _call_llm_stream

logger = logging.getLogger(__name__)
router = APIRouter()


# ---- 请求模型 ----
class ProductItem(BaseModel):
    name: str = ""
    category: str = ""
    price: float | None = None
    platform: str = ""
    sales_volume: int | None = None
    rating: float | None = None
    description: str = ""


class BatchAnalysisRequest(BaseModel):
    products: list[ProductItem]
    deep_thinking: bool = False


class ExportDocxRequest(BaseModel):
    report: str
    products: list[ProductItem] = []
    title: str = "选定商品批量分析报告"


# 批量分析专用系统提示词（多商品综合分析，区别于单品类前景分析）
BATCH_SYSTEM_PROMPT = (
    "你是电商选品分析系统。用户会提供一批已选定的商品数据，请对这些商品进行批量综合分析。"
    "按以下结构输出（板块标题独占一行、前后空行，每个条目以「- 」开头独占一行，条目间空行）：\n\n"
    "一、商品概览\n"
    "- 商品数量与类目分布\n"
    "- 价格区间与平台分布\n"
    "- 销量与评分整体表现\n\n"
    "二、选品优势\n"
    "- 销量口碑表现\n"
    "- 价格竞争力\n"
    "- 类目潜力\n\n"
    "三、潜在风险\n"
    "- 竞争与同质化\n"
    "- 评分与库存隐患\n"
    "- 平台与季节性风险\n\n"
    "四、单品点评\n"
    "- 逐个商品一句话点评（名称 + 核心结论）\n\n"
    "五、行动建议\n"
    "- 优先主推商品\n"
    "- 建议规避或观望商品\n"
    "- 整体选品策略\n\n"
    "规则：正式客观，引用具体数字；禁止口语、感叹号、人称代词；条目间空行。"
)

CN_NUMERALS = ("一、", "二、", "三、", "四、", "五、", "六、", "七、", "八、", "九、", "十、")


def _build_products_text(products: list[ProductItem]) -> str:
    """把选定商品列表拼成给 LLM 的文本。"""
    lines = []
    for i, p in enumerate(products, 1):
        lines.append(
            f"{i}. {p.name} | 类目:{p.category or '无'} | 平台:{p.platform or '无'} | "
            f"价格:{p.price} | 销量:{p.sales_volume} | 评分:{p.rating} | 描述:{p.description or '无'}"
        )
    return "\n".join(lines)


def _sse(data: dict) -> str:
    return f"data: {json.dumps(data, ensure_ascii=False, default=str)}\n\n"


@router.post("/batch")
def batch_analysis(req: BatchAnalysisRequest):
    """对选定商品批量做 AI 分析，流式 SSE 返回 content。"""
    if not req.products:
        return {"detail": "未提供选定商品"}

    products_text = _build_products_text(req.products)
    user_prompt = (
        f"以下是我选定的 {len(req.products)} 件商品数据：\n\n{products_text}\n\n"
        "请对这批商品做批量综合分析，给出选品建议。"
    )
    messages = [
        {"role": "system", "content": BATCH_SYSTEM_PROMPT},
        {"role": "user", "content": user_prompt},
    ]
    model = settings.LLM_REASONER_MODEL if req.deep_thinking else settings.LLM_MODEL

    def event_stream():
        try:
            for kind, text in _call_llm_stream(messages, model=model):
                # 批量分析只转发正文，不展示思考过程
                if kind == "content":
                    yield _sse({"type": "content", "delta": text})
        except Exception as e:
            logger.exception("批量分析流式失败: %s", e)
            yield _sse({"type": "error", "message": "分析失败，请稍后重试"})
        yield _sse({"type": "done"})

    return StreamingResponse(event_stream(), media_type="text/event-stream")


@router.post("/export-docx")
def export_docx(req: ExportDocxRequest):
    """把分析报告导出为 Word（.docx）。python-docx 懒加载，未安装时返回明确错误。"""
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        logger.error("python-docx 未安装，无法导出 Word")
        return {"detail": "服务器未安装 python-docx，无法导出 Word 报告"}

    doc = Document()

    # 主标题
    title = doc.add_heading(req.title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 生成信息
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mrun = meta.add_run(
        f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}　商品数：{len(req.products)}"
    )
    mrun.font.size = Pt(9)
    mrun.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    # 选定商品清单（表格）
    if req.products:
        doc.add_heading("选定商品清单", level=1)
        table = doc.add_table(rows=1, cols=6)
        try:
            table.style = "Light Grid Accent 1"
        except Exception:
            pass
        headers = ["商品名称", "类目", "平台", "价格", "销量", "评分"]
        for i, h in enumerate(headers):
            table.rows[0].cells[i].text = h
        for p in req.products:
            row = table.add_row().cells
            row[0].text = str(p.name)
            row[1].text = str(p.category) if p.category else ""
            row[2].text = str(p.platform) if p.platform else ""
            row[3].text = f"¥{p.price}" if p.price is not None else ""
            row[4].text = str(p.sales_volume) if p.sales_volume is not None else ""
            row[5].text = str(p.rating) if p.rating is not None else ""

    # 报告正文：按行解析
    doc.add_heading("AI 分析报告", level=1)
    for raw_line in req.report.split("\n"):
        line = raw_line.strip()
        if not line:
            continue
        if any(line.startswith(n) for n in CN_NUMERALS):
            doc.add_heading(line, level=2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        else:
            doc.add_paragraph(line)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    # filename 用 ASCII：Starlette 用 latin-1 编码 header，中文会抛 UnicodeEncodeError；
    # 下载时的中文文件名由前端 a.download 指定，后端 filename 仅为 fallback
    filename = f"analysis_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
